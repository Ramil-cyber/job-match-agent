"""Safe, in-memory text extraction for user-supplied documents."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UPLOAD_SIZE_MB = 5
MAX_PDF_PAGES = 40
MAX_DOCX_ARCHIVE_ENTRIES = 1_000
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DOCX_ENTRY_BYTES = 10 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100


class DocumentExtractionError(ValueError):
    """Raised when an uploaded document cannot be accepted safely."""


@dataclass(frozen=True)
class ExtractedDocument:
    """Normalized text and safe metadata from one uploaded document."""

    text: str
    format_name: str

    @property
    def character_count(self) -> int:
        """Return the normalized extracted-text length."""

        return len(self.text)


def _filename_extension(filename: str) -> str:
    """Return a normalized extension without treating the name as a path."""

    safe_name = str(filename).replace("\\", "/").rsplit("/", maxsplit=1)[-1]
    return PurePath(safe_name).suffix.lower()


def _normalize_text(text: str) -> str:
    """Normalize line endings while preserving useful document structure."""

    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = normalized.replace("\x00", "").replace("\x0c", "\n")
    lines = [line.rstrip() for line in normalized.splitlines()]
    normalized = "\n".join(lines)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _validate_extracted_text(text: str, *, max_characters: int) -> str:
    """Reject empty or oversized extraction results."""

    normalized = _normalize_text(text)

    if not normalized or not any(character.isalnum() for character in normalized):
        raise DocumentExtractionError(
            "No readable text was found. If this is a scanned PDF, convert it "
            "with OCR or paste the text instead."
        )
    if len(normalized) > max_characters:
        raise DocumentExtractionError(
            f"The extracted text contains {len(normalized):,} characters; the "
            f"limit is {max_characters:,}. Shorten the document or paste a "
            "shorter version."
        )

    return normalized


def _extract_txt(data: bytes) -> str:
    """Decode a UTF-8 text file, including an optional byte-order mark."""

    if b"\x00" in data:
        raise DocumentExtractionError(
            "The TXT file does not appear to be UTF-8 plain text."
        )

    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentExtractionError(
            "The TXT file must use UTF-8 text encoding."
        ) from error


def _extract_pdf(data: bytes) -> str:
    """Extract text from a bounded, unencrypted PDF."""

    if b"%PDF-" not in data[:1_024]:
        raise DocumentExtractionError(
            "The file extension is PDF, but the file content is not a valid PDF."
        )

    try:
        reader = PdfReader(BytesIO(data), strict=False)
    except (PdfReadError, OSError, TypeError, ValueError) as error:
        raise DocumentExtractionError(
            "The PDF is damaged or could not be read."
        ) from error

    if reader.is_encrypted:
        raise DocumentExtractionError(
            "Password-protected or encrypted PDFs are not supported."
        )
    if not reader.pages:
        raise DocumentExtractionError("The PDF does not contain any pages.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentExtractionError(
            f"The PDF has too many pages. The limit is {MAX_PDF_PAGES}."
        )

    page_text: list[str] = []
    try:
        for page in reader.pages:
            page_text.append(page.extract_text() or "")
    except (MemoryError, PdfReadError, OSError, TypeError, ValueError) as error:
        raise DocumentExtractionError(
            "The PDF text could not be extracted safely."
        ) from error

    return "\n\n".join(page_text)


def _validate_docx_archive(data: bytes) -> None:
    """Reject malformed, macro-enabled, or suspicious DOCX archives."""

    if not data.startswith(b"PK"):
        raise DocumentExtractionError(
            "The file extension is DOCX, but the file content is not a valid "
            "DOCX document."
        )

    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}

            if len(entries) > MAX_DOCX_ARCHIVE_ENTRIES:
                raise DocumentExtractionError(
                    "The DOCX archive contains too many internal files."
                )
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                raise DocumentExtractionError(
                    "The uploaded file is not a complete DOCX document."
                )
            if any(
                name.lower().endswith("vbaproject.bin")
                or name.lower().endswith("vbadata.xml")
                for name in names
            ):
                raise DocumentExtractionError(
                    "Macro-enabled Word documents are not supported."
                )

            total_uncompressed = sum(entry.file_size for entry in entries)
            if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentExtractionError(
                    "The expanded DOCX document is too large to process safely."
                )

            for entry in entries:
                if entry.file_size > MAX_DOCX_ENTRY_BYTES:
                    raise DocumentExtractionError(
                        "The DOCX document contains an oversized internal file."
                    )
                if (
                    entry.compress_size > 0
                    and entry.file_size / entry.compress_size
                    > MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise DocumentExtractionError(
                        "The DOCX compression ratio is too high to process safely."
                    )
    except BadZipFile as error:
        raise DocumentExtractionError(
            "The DOCX document is damaged or could not be read."
        ) from error


def _table_lines(table: Table) -> list[str]:
    """Extract table cells row by row without repeating merged cells."""

    lines: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        seen_cells: set[int] = set()
        for cell in row.cells:
            cell_identity = id(cell._tc)
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)
            cell_text = _normalize_text("\n".join(p.text for p in cell.paragraphs))
            if cell_text:
                cells.append(cell_text.replace("\n", " "))
        if cells:
            lines.append(" | ".join(cells))
    return lines


def _extract_docx(data: bytes) -> str:
    """Extract paragraphs and table cells from a DOCX in document order."""

    _validate_docx_archive(data)

    try:
        document = Document(BytesIO(data))
        blocks: list[str] = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                if block.text.strip():
                    blocks.append(block.text)
            elif isinstance(block, Table):
                blocks.extend(_table_lines(block))
        return "\n".join(blocks)
    except (BadZipFile, KeyError, OSError, TypeError, ValueError) as error:
        raise DocumentExtractionError(
            "The DOCX document is damaged or could not be read."
        ) from error


def extract_document(
    data: bytes,
    filename: str,
    *,
    max_characters: int,
    max_file_bytes: int = MAX_UPLOAD_BYTES,
) -> ExtractedDocument:
    """Validate and extract one supported upload without writing it to disk."""

    if not isinstance(data, bytes):
        raise TypeError("Document data must be bytes.")
    if max_characters < 1 or max_file_bytes < 1:
        raise ValueError("Document limits must be positive.")
    if not data:
        raise DocumentExtractionError("The uploaded file is empty.")
    if len(data) > max_file_bytes:
        raise DocumentExtractionError(
            f"The uploaded file is larger than the {max_file_bytes / 1_048_576:g} "
            "MB limit."
        )

    extension = _filename_extension(filename)
    if extension == ".doc":
        raise DocumentExtractionError(
            "Legacy .doc files are not supported. Save the file as DOCX, PDF, "
            "or UTF-8 TXT."
        )
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            "Unsupported file type. Upload a PDF, DOCX, or UTF-8 TXT file."
        )

    try:
        if extension == ".pdf":
            text = _extract_pdf(data)
            format_name = "PDF"
        elif extension == ".docx":
            text = _extract_docx(data)
            format_name = "DOCX"
        else:
            text = _extract_txt(data)
            format_name = "TXT"
    except DocumentExtractionError:
        raise
    except Exception as error:  # noqa: BLE001
        raise DocumentExtractionError(
            "The document could not be processed safely."
        ) from error

    return ExtractedDocument(
        text=_validate_extracted_text(text, max_characters=max_characters),
        format_name=format_name,
    )
