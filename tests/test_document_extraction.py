import unittest
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from job_match_agent.document_extraction import (
    MAX_PDF_PAGES,
    DocumentExtractionError,
    extract_document,
)


def build_pdf(text: str | None = None) -> bytes:
    """Create a small in-memory PDF for extraction tests."""

    output = BytesIO()
    pdf = canvas.Canvas(output)
    if text:
        pdf.drawString(72, 720, text)
    pdf.showPage()
    pdf.save()
    return output.getvalue()


def build_docx() -> bytes:
    """Create a DOCX containing both paragraphs and table cells."""

    output = BytesIO()
    document = Document()
    document.add_paragraph("Jordan Example")
    document.add_paragraph("Python and SQL analyst")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Evidence"
    table.cell(1, 0).text = "Forecasting"
    table.cell(1, 1).text = "Reduced error by 12 percent"
    document.save(output)
    return output.getvalue()


class DocumentExtractionTests(unittest.TestCase):
    def test_utf8_txt_is_normalized(self):
        result = extract_document(
            b"Resume\r\n\r\n\r\nPython and SQL\r\n",
            "resume.txt",
            max_characters=10_000,
        )

        self.assertEqual(result.format_name, "TXT")
        self.assertEqual(result.text, "Resume\n\nPython and SQL")
        self.assertEqual(result.character_count, len(result.text))

    def test_utf8_bom_txt_is_supported(self):
        result = extract_document(
            b"\xef\xbb\xbfJob description with analytics experience",
            "job.txt",
            max_characters=10_000,
        )

        self.assertTrue(result.text.startswith("Job description"))

    def test_non_utf8_txt_is_rejected(self):
        with self.assertRaisesRegex(DocumentExtractionError, "UTF-8"):
            extract_document(
                b"Resume \xff",
                "resume.txt",
                max_characters=10_000,
            )

    def test_text_pdf_is_extracted(self):
        result = extract_document(
            build_pdf("Python SQL forecasting experience"),
            "resume.pdf",
            max_characters=10_000,
        )

        self.assertEqual(result.format_name, "PDF")
        self.assertIn("Python SQL forecasting experience", result.text)

    def test_scanned_or_image_only_pdf_is_rejected(self):
        with self.assertRaisesRegex(DocumentExtractionError, "scanned PDF"):
            extract_document(
                build_pdf(),
                "resume.pdf",
                max_characters=10_000,
            )

    def test_encrypted_pdf_is_rejected(self):
        reader = PdfReader(BytesIO(build_pdf("Private resume")))
        writer = PdfWriter()
        writer.append_pages_from_reader(reader)
        writer.encrypt("password")
        output = BytesIO()
        writer.write(output)

        with self.assertRaisesRegex(DocumentExtractionError, "encrypted PDFs"):
            extract_document(
                output.getvalue(),
                "resume.pdf",
                max_characters=10_000,
            )

    def test_pdf_page_limit_is_enforced(self):
        writer = PdfWriter()
        for _ in range(MAX_PDF_PAGES + 1):
            writer.add_blank_page(width=612, height=792)
        output = BytesIO()
        writer.write(output)

        with self.assertRaisesRegex(DocumentExtractionError, "too many pages"):
            extract_document(
                output.getvalue(),
                "resume.pdf",
                max_characters=10_000,
            )

    def test_docx_paragraphs_and_tables_are_extracted(self):
        result = extract_document(
            build_docx(),
            "resume.docx",
            max_characters=10_000,
        )

        self.assertEqual(result.format_name, "DOCX")
        self.assertIn("Jordan Example", result.text)
        self.assertIn("Skill | Evidence", result.text)
        self.assertIn("Forecasting | Reduced error by 12 percent", result.text)

    def test_macro_enabled_docx_is_rejected(self):
        output = BytesIO(build_docx())
        with ZipFile(output, mode="a", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/vbaProject.bin", b"macro")

        with self.assertRaisesRegex(DocumentExtractionError, "Macro-enabled"):
            extract_document(
                output.getvalue(),
                "resume.docx",
                max_characters=10_000,
            )

    def test_suspicious_docx_expansion_is_rejected(self):
        output = BytesIO(build_docx())
        with ZipFile(output, mode="a", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/media/oversized.bin", b"0" * (21 * 1024 * 1024))

        with self.assertRaisesRegex(DocumentExtractionError, "expanded DOCX"):
            extract_document(
                output.getvalue(),
                "resume.docx",
                max_characters=10_000,
            )

    def test_mismatched_pdf_content_is_rejected(self):
        with self.assertRaisesRegex(DocumentExtractionError, "not a valid PDF"):
            extract_document(
                b"This is plain text, not a PDF.",
                "resume.pdf",
                max_characters=10_000,
            )

    def test_mismatched_docx_content_is_rejected(self):
        with self.assertRaisesRegex(DocumentExtractionError, "not a valid DOCX"):
            extract_document(
                b"This is plain text, not a DOCX.",
                "resume.docx",
                max_characters=10_000,
            )

    def test_legacy_doc_is_rejected(self):
        with self.assertRaisesRegex(DocumentExtractionError, "Legacy .doc"):
            extract_document(
                b"legacy document",
                "resume.doc",
                max_characters=10_000,
            )

    def test_unsupported_extension_is_rejected(self):
        with self.assertRaisesRegex(DocumentExtractionError, "Unsupported file type"):
            extract_document(
                b"resume",
                "resume.rtf",
                max_characters=10_000,
            )

    def test_file_size_limit_is_enforced_before_parsing(self):
        with self.assertRaisesRegex(DocumentExtractionError, "larger than"):
            extract_document(
                b"0123456789",
                "resume.txt",
                max_characters=10_000,
                max_file_bytes=5,
            )

    def test_extracted_character_limit_is_enforced(self):
        with self.assertRaisesRegex(DocumentExtractionError, "limit is 10"):
            extract_document(
                b"12345678901",
                "resume.txt",
                max_characters=10,
            )

    def test_filename_is_not_used_as_a_filesystem_path(self):
        result = extract_document(
            b"Safe resume text",
            "../../private/resume.txt",
            max_characters=10_000,
        )

        self.assertEqual(result.text, "Safe resume text")


if __name__ == "__main__":
    unittest.main()
